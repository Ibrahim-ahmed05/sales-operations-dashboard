from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\Administrator\Desktop\Sales Dashboard\pixel-perfect-replication\Meridian Dashboard Implementation Overview.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table, color="D9E2F0"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tbl_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tbl_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="172239", size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Normal"].font.color.rgb = RGBColor(45, 58, 80)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.12
for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 12)):
    style = styles[name]
    style.font.name = "Aptos Display" if name == "Title" else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(6)

title = doc.add_paragraph(style="Title")
title.add_run("Meridian Dashboard Implementation Overview")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run("A simple guide to what was built and how it works")
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(75, 105, 160)

doc.add_paragraph(
    "This document explains the completed Meridian sales and operations dashboard in plain language. "
    "It covers the pages, KPIs, data rules, user experience, and the checks completed before delivery. "
    "The KPI definitions and API requirements used for the implementation were provided by Haroon bhai."
)

doc.add_heading("What the dashboard does", level=1)
doc.add_paragraph(
    "Meridian turns the supplied sales, orders, inventory, returns, customers, products, invoices, and payments "
    "into a single management view. A user can move between focused pages, change the date range, inspect charts, "
    "open source-record details, and export the visible records. The dashboard opens directly in a demo workspace "
    "so no login screen is required for the concept presentation."
)
add_bullets(doc, [
    "The imported source contains 3,200 orders, 11,486 order lines, 2,369 invoices, 120 customers, and 72 products, along with inventory, returns, payments, and stock movements.",
    "The dashboard uses the full imported order period as its initial view, so the first screen shows meaningful sales immediately.",
    "The interface uses the Meridian brand, a clean blue visual system, professional typography, restrained shadows, and clear card hierarchy.",
])

doc.add_heading("Dashboard pages and KPIs", level=1)
table = doc.add_table(rows=1, cols=3)
table.autofit = True
borders(table)
headers = ["Page", "Main KPIs", "Supporting information"]
for i, text in enumerate(headers):
    set_cell_text(table.rows[0].cells[i], text, bold=True, color="FFFFFF", size=9)
    shade(table.rows[0].cells[i], "285BC7")
rows = [
    ("Sales", "Total sales, total orders, average order value, sales growth", "Sales over time, category sales, period comparison, order value spread"),
    ("Orders", "Pending, delivered, cancelled orders, fulfilment rate", "Status breakdown, monthly orders, estimated time in each status, fulfilment trend"),
    ("Inventory", "Low-stock items, out-of-stock items, active products, stock value", "Low-stock list, stock movement, fastest-moving products, category stock value"),
    ("Receivables", "Outstanding amount, overdue amount, overdue invoices, average days to pay", "Aging bands, collection trend, invoice status, top overdue customers"),
    ("Top performers", "Top product and top customer", "Product ranking, customer ranking, customer segment mix"),
    ("Operations", "Average fulfilment time, return rate, repeat customer rate, inventory turnover", "Fulfilment trend, return trend, repeat customers, turnover trend"),
]
for row in rows:
    cells = table.add_row().cells
    for i, text in enumerate(row):
        set_cell_text(cells[i], text, size=8.5)
    if len(table.rows) % 2 == 0:
        for cell in cells:
            shade(cell, "F5F8FC")

doc.add_heading("How the numbers are handled", level=1)
doc.add_paragraph(
    "The dashboard follows the agreed business definitions so that the same number is used consistently in cards, "
    "charts, and drill-downs. Cancelled orders are excluded from sales and fulfilment calculations. Sales are based "
    "on line quantity multiplied by sale price, with returns reflected where applicable. Outstanding receivables are "
    "invoiced amounts minus payments, and overdue balances use the dashboard reference date. Available inventory is "
    "on-hand stock minus reserved stock where that view is used."
)
add_bullets(doc, [
    "Date filters are inclusive and update the selected page, KPI cards, charts, comparisons, and record details.",
    "Inventory counts and receivables balances are treated as current snapshots when the metric definition calls for a snapshot.",
    "Charts use dynamic axis steps so labels remain evenly spaced and readable across different value ranges.",
    "Charts animate gently when loaded, with reduced-motion support for users who prefer minimal animation.",
])

doc.add_heading("Illustrative concept values", level=1)
doc.add_paragraph(
    "The supplied dataset is a fabricated concept dataset, so a small number of operational fields are not present in "
    "the source. To keep the dashboard visually complete for presentation, these fields use clearly labeled illustrative "
    "estimates rather than blank cards. This applies to average time in each status, inventory turnover when historical "
    "inventory value cannot be calculated, average days to pay when no paid invoice is in the selected period, and sales "
    "growth when no prior-period records exist. The descriptions identify these values as estimates."
)

doc.add_heading("Responsive and visual improvements", level=1)
doc.add_paragraph(
    "The dashboard was adjusted for desktop, laptop, tablet, and mobile widths. On compact screens, navigation wraps, "
    "KPI cards reflow into smaller columns, and chart panels stack vertically. When the available height is too small "
    "for every module, the page can scroll naturally instead of cropping content. Chart labels are anchored inside the "
    "card edges, and panels contain their graph artwork so values do not spill outside the card."
)

doc.add_heading("Traceability and interactions", level=1)
add_bullets(doc, [
    "Most cards, charts, rankings, and tables can open a detail drawer with the source records behind the number.",
    "Detail views support searching, pagination, record-to-record drill-down, and CSV export for the visible page.",
    "Metric definitions are available inside the dashboard so a CEO or manager can understand the reporting basis without leaving the page.",
    "The dashboard uses a responsive layout and keeps the most important information visible without unnecessary decoration.",
])

doc.add_heading("Validation completed", level=1)
doc.add_paragraph(
    "The implementation was checked through TypeScript compilation, API and metric verification, responsive layout review, "
    "chart rendering checks, and a production build. The API verification suite completed 86 metric, API, drill-down, "
    "date, cache, and authorization checks successfully."
)

doc.add_heading("Acknowledgement", level=1)
doc.add_paragraph(
    "The KPI definitions, page requirements, and API information used for this dashboard were provided by Haroon bhai."
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Meridian Dashboard · Implementation overview")
footer_run.font.name = "Aptos"
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(120, 135, 160)

doc.save(OUT)
print(OUT)
